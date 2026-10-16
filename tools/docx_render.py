"""Minimal Markdown -> .docx renderer for BRDs (headings, paragraphs, lists, tables, bold/italic/code)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt

_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _add_inline(paragraph, text: str) -> None:
    for chunk in _INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            paragraph.add_run(chunk[1:-1]).italic = True
        else:
            paragraph.add_run(chunk)


def _cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def _flush_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row in enumerate(rows):
        for j in range(width):
            cell = table.cell(i, j)
            cell.text = ""
            para = cell.paragraphs[0]
            _add_inline(para, row[j] if j < len(row) else "")
            if i == 0:
                for run in para.runs:
                    run.bold = True
    doc.add_paragraph()


def markdown_to_docx(markdown: str, path: Path) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    table_rows: list[list[str]] = []
    for raw in markdown.splitlines():
        line = raw.rstrip()

        if line.strip().startswith("|"):
            if not _TABLE_SEP.match(line):
                table_rows.append(_cells(line))
            continue
        if table_rows:
            _flush_table(doc, table_rows)
            table_rows = []

        stripped = line.strip()
        if not stripped or stripped in {"---", "***"}:
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            doc.add_heading(heading.group(2).strip(), level=0 if level == 1 else level - 1)
            continue
        bullet = re.match(r"^(\s*)[-*+]\s+(?:\[[ xX]\]\s+)?(.*)$", line)
        if bullet:
            depth = len(bullet.group(1)) // 2
            para = doc.add_paragraph(style="List Bullet 2" if depth else "List Bullet")
            _add_inline(para, bullet.group(2))
            continue
        numbered = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if numbered:
            _add_inline(doc.add_paragraph(style="List Number"), numbered.group(1))
            continue
        _add_inline(doc.add_paragraph(), stripped)

    if table_rows:
        _flush_table(doc, table_rows)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
