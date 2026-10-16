"""Offline tests for BRD_Agent_V2: stores, service flow, chat agent. No GCP or model calls.

    cd BRD_Agent_V2 && pytest -q
"""
from __future__ import annotations

import io
import json
import re
import sys
from pathlib import Path
from typing import AsyncGenerator

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT.parent))

from docx import Document  # noqa: E402
from google.adk.models.base_llm import BaseLlm  # noqa: E402
from google.adk.models.llm_request import LlmRequest  # noqa: E402
from google.adk.models.llm_response import LlmResponse  # noqa: E402
from google.genai import types  # noqa: E402

MEETING = "MTG-2026-10-15-INVOICE-ALERTS"
SAMPLE = ROOT / "samples" / "invoice_approval_alerts_call.txt"
SCHEMA = {"meeting_id", "summary", "MOM", "jira_actions_requirend", "jira_action_description", "bdr_generated",
          "relevant_jira", "user_confirmation_provided", "last_modified_date", "last_modified_time_stamp"}

INSIGHTS = {
    "meeting_title": "Invoice Approval Alerts - Requirements Discussion", "meeting_date": "2026-10-15",
    "facilitator": "Priya Nair",
    "attendees": [{"name": n, "role": None, "joined_late": False} for n in
                  ["Priya Nair", "Rohan Mehta", "Anita Desai", "Karthik Iyer", "Sneha Rao", "Vikram Shetty"]],
    "purpose": "p", "problem_statement": "p", "current_state": "c", "discussion_summary": ["SLA"],
    "requirements": [], "business_rules": [], "decisions": [], "out_of_scope": [], "dependencies": [],
    "risks": [], "success_metrics": [], "timeline": None, "action_items": [], "open_questions": [],
    "mom_extra_cc": ["Meera"], "jira_requests": ["Create epic (13:52)"],
}
PLAN = {"project_key": "FINAUTO", "notes": ["Dashboard parked"], "actions": [
    {"ref": "A1", "action": "create_issue", "issue_type": "Epic", "summary": "Invoice approval SLA alerts",
     "rationale": "asked at 13:52", "source_timestamp": "13:52"},
    {"ref": "A2", "action": "create_issue", "issue_type": "Story", "summary": "Bundled Teams reminders",
     "parent": "A1", "assignee": "Karthik", "acceptance_criteria": ["Given x When y Then z"], "rationale": "14:10"},
    {"ref": "A3", "action": "link_issues", "link_type": "Duplicate", "issue": "FINAUTO-1098",
     "other_issue": "A1", "rationale": "13:52"},
    {"ref": "A4", "action": "link_issues", "link_type": "Relates", "issue": "1142", "other_issue": "A1",
     "rationale": "14:05"},
]}
SUMMARY = "The team scoped SLA-based reminders for stalled invoice approvals.\n\n**Key points**\n- Teams first"
MOM = "Hi all,\n\n### Decisions\n- Teams first, email fallback\n\n| # | Action | Owner |\n|  ---  |  ---  |  ---  |"


def _between(text: str, start: str, end: str) -> str:
    return text.split(start, 1)[1].split(end, 1)[0] if start in text else ""


class ScriptedLlm(BaseLlm):
    model: str = "scripted"

    async def generate_content_async(self, llm_request: LlmRequest, stream: bool = False
                                     ) -> AsyncGenerator[LlmResponse, None]:
        system = str(llm_request.config.system_instruction or "")
        if "FILLING IN THE TEMPLATE" in system:
            tpl = _between(system, "TEMPLATE:\n<<<\n", "\n>>>")
            heads = [l for l in tpl.splitlines() if l.startswith("#")]
            text = "\n\n".join(f"{h}\nContent for this section, grounded in the meeting. " * 1 for h in heads)
        elif system.startswith("You are HeyPM"):
            text, call = self._chat(llm_request)
            if call:
                yield LlmResponse(content=types.Content(role="model", parts=[types.Part(function_call=call)]))
                return
        elif 'You write the "AI Summary"' in system:
            text = SUMMARY
        elif "Minutes of Meeting" in system:
            text = MOM
        elif "JIRA action plan" in system:
            text = json.dumps(PLAN)
        elif "extract structured meeting insights" in system:
            text = json.dumps(INSIGHTS)
        else:
            text = "?"
        yield LlmResponse(content=types.Content(role="model", parts=[types.Part(text=text)]))

    @staticmethod
    def _chat(req: LlmRequest):
        user_texts = [p.text for c in req.contents if c.role == "user" for p in (c.parts or []) if p.text]
        responses = [p.function_response for c in req.contents for p in (c.parts or []) if p.function_response]
        last = user_texts[-1] if user_texts else ""
        done = [r for r in responses if r.name == "confirm_jira_actions"
                and (r.response or {}).get("status") != "pending_approval"]
        if "create the jira" in last.lower():
            if done:
                return f"Done: {done[-1].response.get('relevant_jira')}", None
            return "", types.FunctionCall(name="confirm_jira_actions",
                                          args={"skip_refs": [], "send_mom_email": True})
        if "brd" in last.lower():
            gen = [r for r in responses if r.name == "generate_brd"]
            if gen:
                return f"BRD ready: {gen[-1].response.get('docx_uri')}", None
            return "", types.FunctionCall(name="generate_brd", args={"template": ""})
        has_transcript = "Anita Desai" in str(req.config.system_instruction)
        return ("Anita said the SLA is three working days (3:22)." if has_transcript else "No meeting open."), None


@pytest.fixture()
def svc(tmp_path, monkeypatch):
    monkeypatch.setenv("STORAGE_BACKEND", "local")
    monkeypatch.setenv("DRY_RUN", "true")
    monkeypatch.setenv("LOCAL_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("PEOPLE_FILE", str(ROOT / "people.yaml"))
    (tmp_path / "transcripts").mkdir()
    (tmp_path / "transcripts" / f"{MEETING}.txt").write_text(SAMPLE.read_text(encoding="utf-8"), encoding="utf-8")
    from BRD_Agent_V2 import agent, pipeline, service
    fake = ScriptedLlm()
    for a in pipeline.LLM_AGENTS + [agent.root_agent]:
        a.model = fake
    s = service.HeyPMService()
    service.set_service(s)
    yield s
    service.set_service(None)


# ---------------------------------------------------------------------------
# Stores
# ---------------------------------------------------------------------------
def test_vtt_and_json_transcripts_normalise():
    from BRD_Agent_V2.stores.transcripts import json_to_text, vtt_to_text
    vtt = ("WEBVTT\n\n1\n00:03:22.000 --> 00:03:25.000\n<v Anita Desai>Normal invoices, three working days.</v>\n\n"
           "2\n00:03:25.100 --> 00:03:27.000\n<v Anita Desai>Discount ones, one day.</v>\n")
    assert vtt_to_text(vtt) == "Anita Desai   3:22\nNormal invoices, three working days. Discount ones, one day."
    js = json.dumps({"segments": [{"speaker": "Rohan", "start": 95, "text": "Teams first."}]})
    assert json_to_text(js) == "Rohan   1:35\nTeams first."


def test_templates_default_custom_and_docx(tmp_path):
    from BRD_Agent_V2.stores.templates import TemplateStore
    d = tmp_path / "tpl"
    d.mkdir()
    assert TemplateStore(local_dir=d).get(None).id == "default"
    doc = Document()
    doc.add_heading("Finance BRD", 0)
    doc.add_heading("Purpose", 1)
    doc.add_paragraph("Why we are doing this.")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Control", "Owner"
    doc.add_heading("Controls Impact", 1)
    buf = io.BytesIO()
    doc.save(buf)
    (d / "default.docx").write_bytes(buf.getvalue())
    store = TemplateStore(local_dir=d)
    tpl = store.get(None)
    assert tpl.name == "default" and "## Controls Impact" in tpl.content and "| Control | Owner |" in tpl.content
    assert store.get("default").id == "default"  # explicit 'default' = bundled


def test_firestore_store_keeps_schema_top_level():
    from BRD_Agent_V2.stores.meetings import FirestoreMeetingStore, blank_record

    class Snap:
        def __init__(self, d): self._d, self.exists = d, d is not None
        def to_dict(self): return json.loads(json.dumps(self._d))

    class Doc:
        def __init__(self, db, k): self.db, self.k = db, k
        def get(self): return Snap(self.db.get(self.k))
        def set(self, d): self.db[self.k] = d

    class Col:
        def __init__(self): self.db = {}
        def document(self, k): return Doc(self.db, k)

    class Client:
        def __init__(self): self.c = Col()
        def collection(self, name): return self.c

    store = FirestoreMeetingStore(client=Client(), collection="meetings")
    store.put(blank_record("M1"))
    rec = store.update("M1", {"summary": "s", "agent_data": {"status": "processed"}})
    assert set(rec) == SCHEMA | {"agent_data"} and rec["summary"] == "s"
    assert re.match(r"\d{4}-\d{2}-\d{2}$", rec["last_modified_date"])
    assert "T" in rec["last_modified_time_stamp"]


# ---------------------------------------------------------------------------
# Service flow
# ---------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_process_confirm_and_brd_flow(svc, tmp_path):
    from BRD_Agent_V2.service import AlreadyConfirmed

    rec = await svc.process_meeting(MEETING)
    assert set(rec) == SCHEMA
    assert rec["summary"].startswith("The team scoped")
    assert "|  ---" not in rec["MOM"]  # padded table cleaned
    assert rec["jira_actions_requirend"] is True
    assert "A1" in rec["jira_action_description"] and "Dashboard parked" in rec["jira_action_description"]
    assert rec["relevant_jira"] == ["FINAUTO-1098", "FINAUTO-1142"]
    assert rec["bdr_generated"] is False and rec["user_confirmation_provided"] is False

    full = await svc.confirm_actions(MEETING, skip_refs=["A4"])
    assert full["user_confirmation_provided"] is True
    assert full["agent_data"]["mom_email"]["status"] == "dry_run"
    assert "meera.kulkarni@example.com" in full["agent_data"]["mom_email"]["cc"]
    assert full["agent_data"]["jira_results"]["A4"]["status"] == "skipped"
    assert "FINAUTO-NEW1" in full["relevant_jira"] and "FINAUTO-NEW2" in full["relevant_jira"]

    again = await svc.confirm_actions(MEETING)  # idempotent: nothing re-created, A4 now runs
    assert again["agent_data"]["jira_results"]["A1"]["key"] == "FINAUTO-NEW1"
    assert again["agent_data"]["jira_results"]["A4"]["status"] == "ok"

    with pytest.raises(AlreadyConfirmed):
        await svc.process_meeting(MEETING)

    # BRD only on request, using a custom default template
    (tmp_path / "templates").mkdir()
    (tmp_path / "templates" / "default.md").write_text("# Finance BRD\n## Purpose\n## Controls Impact\n",
                                                       encoding="utf-8")
    svc.stores.templates.local_dir = tmp_path / "templates"
    brd = await svc.generate_brd(MEETING)
    assert brd["bdr_generated"] is True
    info = brd["agent_data"]["brd"]
    assert info["template_name"] == "default" and info["warnings"] == []
    text = "\n".join(p.text for p in Document(info["docx_uri"]).paragraphs)
    assert "Controls Impact" in text
    assert svc.get_record(MEETING)["bdr_generated"] is True


@pytest.mark.asyncio
async def test_generate_brd_processes_first_and_uses_bundled_default(svc):
    rec = await svc.generate_brd(MEETING, template_id="default")
    assert rec["bdr_generated"] and rec["summary"] and rec["agent_data"]["brd"]["template_id"] == "default"
    assert Path(rec["agent_data"]["brd"]["md_uri"]).read_text(encoding="utf-8").startswith("# Business Requirements")


# ---------------------------------------------------------------------------
# Chat agent
# ---------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_chat_answers_and_confirms_actions(svc):
    await svc.process_meeting(MEETING)
    r = await svc.ask(MEETING, "What did Anita say about the SLA?")
    assert "three working days" in r["answer"] and r["pending_confirmation"] is None

    r2 = await svc.ask(MEETING, "Please create the JIRA stories", session_id=r["session_id"])
    pending = r2["pending_confirmation"]
    assert pending and "Create Epic" in pending["hint"]

    r3 = await svc.answer_confirmation(r["session_id"], pending["confirmation_id"], True, {"skip_refs": []})
    assert "FINAUTO-NEW1" in r3["answer"]
    assert svc.get_record(MEETING)["user_confirmation_provided"] is True


@pytest.mark.asyncio
async def test_chat_generates_brd_on_request(svc):
    await svc.process_meeting(MEETING)
    r = await svc.ask(MEETING, "Generate the BRD please")
    assert "BRD ready" in r["answer"]
    assert svc.get_record(MEETING)["bdr_generated"] is True


# ---------------------------------------------------------------------------
# GCS paths (fake storage client)
# ---------------------------------------------------------------------------
class FakeBlob:
    def __init__(self, store, bucket, name): self.store, self.bucket, self.name = store, bucket, name
    def exists(self): return (self.bucket, self.name) in self.store
    def download_as_bytes(self): return self.store[(self.bucket, self.name)]
    def upload_from_string(self, data, content_type=None):
        self.store[(self.bucket, self.name)] = data if isinstance(data, bytes) else data.encode()


class FakeBucket:
    def __init__(self, store, name): self.store, self.name = store, name
    def blob(self, name): return FakeBlob(self.store, self.name, name)


class FakeGcs:
    def __init__(self): self.store = {}
    def bucket(self, name): return FakeBucket(self.store, name)
    def list_blobs(self, bucket, prefix=None):
        return [FakeBlob(self.store, b, n) for (b, n) in sorted(self.store) if b == bucket
                and (prefix is None or n.startswith(prefix))]


def test_gcs_transcripts_templates_and_outputs(monkeypatch):
    monkeypatch.setenv("TRANSCRIPT_BUCKET", "heypm-data")
    monkeypatch.setenv("TRANSCRIPT_PREFIX", "transcripts")
    monkeypatch.setenv("TEMPLATE_PREFIX", "templates/brd")
    monkeypatch.setenv("OUTPUT_PREFIX", "brds")
    from BRD_Agent_V2.stores.files import GcsFileStore
    from BRD_Agent_V2.stores.templates import TemplateStore
    from BRD_Agent_V2.stores.transcripts import GcsTranscriptStore, TranscriptNotFound

    gcs = FakeGcs()
    gcs.bucket("heypm-data").blob("transcripts/M1.vtt").upload_from_string(
        "WEBVTT\n\n00:00:04.000 --> 00:00:06.000\n<v Priya Nair>Hi all.</v>\n")
    gcs.bucket("heypm-data").blob("transcripts/M2/recording-transcript.json").upload_from_string(
        json.dumps([{"speaker": "Rohan", "start": 12, "text": "Bye."}]))
    t = GcsTranscriptStore(client=gcs)
    assert t.get("M1").text == "Priya Nair   0:04\nHi all." and t.get("M1").source_uri.endswith("M1.vtt")
    assert t.get("M2").text == "Rohan   0:12\nBye."
    assert t.get("gs://heypm-data/transcripts/M1.vtt").text.startswith("Priya")
    with pytest.raises(TranscriptNotFound):
        t.get("M3")

    gcs.bucket("heypm-data").blob("templates/brd/Finance BRD.md").upload_from_string("# Finance\n## Controls\n")
    store = TemplateStore(gcs_client=gcs, source="gcs")
    assert [x["name"] if isinstance(x, dict) else x.name for x in store.list()] == ["Finance BRD",
                                                                                   "Default BRD template"]
    assert store.get("finance brd").content.endswith("## Controls")
    assert store.get(None).id == "default"  # no user default -> bundled

    uri = GcsFileStore(client=gcs).save("M1/BRD.md", b"x", "text/markdown")
    assert uri == "gs://heypm-data/brds/M1/BRD.md" and ("heypm-data", "brds/M1/BRD.md") in gcs.store
