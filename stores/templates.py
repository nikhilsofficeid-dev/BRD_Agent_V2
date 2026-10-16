"""BRD templates from the Templates tab.

Where templates can live (TEMPLATE_SOURCE=auto checks all, in this order):
  - Firestore collection FIRESTORE_TEMPLATES_COLLECTION (default "brd_templates"), documents like
      {"name": "Finance BRD", "content": "<markdown>", "is_default": false}
    or {"name": ..., "gcs_uri": "gs://bucket/path/template.docx", "is_default": ...}
  - GCS folder gs://TEMPLATE_BUCKET/TEMPLATE_PREFIX/  (.docx, .md, .txt files; a file named default.* is the default)
  - The bundled default: templates/default_brd_template.md (always available)

Template ids: "default", "firestore:<doc id>", "gcs:<object name>". The service also accepts a bare name.
A .docx template is converted to a Markdown skeleton (headings, guidance text, table headers).
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from .. import config

TEMPLATE_EXTS = (".docx", ".md", ".txt")


@dataclass
class Template:
    id: str
    name: str
    source: str
    content: str = ""  # markdown skeleton
    is_default: bool = False
    meta: dict = field(default_factory=dict)


class TemplateNotFound(LookupError):
    pass


def docx_to_markdown(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    out: list[str] = []
    body = doc.element.body
    tables = {t._tbl: t for t in doc.tables}
    paragraphs = {p._p: p for p in doc.paragraphs}
    for child in body.iterchildren():
        if child in paragraphs:
            p = paragraphs[child]
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style == "title":
                out.append(f"# {text}")
            elif style.startswith("heading"):
                try:
                    level = int(style.split()[-1])
                except ValueError:
                    level = 1
                out.append(f"{'#' * min(level + 1, 4)} {text}")
            elif "list" in style:
                out.append(f"- {text}")
            else:
                out.append(text)
        elif child in tables:
            rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in tables[child].rows]
            if rows:
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("|" + "---|" * len(rows[0]))
                for r in rows[1:]:
                    out.append("| " + " | ".join(r) + " |")
        out.append("")
    return "\n".join(out).strip()


def _decode(name: str, data: bytes) -> str:
    if name.lower().endswith(".docx"):
        return docx_to_markdown(data)
    return data.decode("utf-8", errors="replace").strip()


def bundled_default() -> Template:
    path = config.default_template_path()
    return Template(id="default", name="Default BRD template", source="bundled",
                    content=path.read_text(encoding="utf-8"), is_default=True)


class TemplateStore:
    """Combines the configured sources. Any backend can be None (e.g. local dev without GCP)."""

    def __init__(self, firestore_client=None, gcs_client=None, local_dir: Optional[Path] = None,
                 source: Optional[str] = None):
        self.source = source or config.template_source()
        self.fs = firestore_client
        self.gcs = gcs_client
        self.local_dir = local_dir

    # -- listing -----------------------------------------------------------------
    def _firestore_templates(self) -> list[Template]:
        if self.fs is None or self.source not in {"auto", "firestore"}:
            return []
        out = []
        for snap in self.fs.collection(config.templates_collection()).stream():
            d = snap.to_dict() or {}
            out.append(Template(id=f"firestore:{snap.id}", name=d.get("name") or snap.id, source="firestore",
                                is_default=bool(d.get("is_default")), meta=d))
        return out

    def _gcs_templates(self) -> list[Template]:
        if self.gcs is None or self.source not in {"auto", "gcs"} or not config.template_bucket():
            return []
        prefix = config.template_prefix()
        out = []
        for blob in self.gcs.list_blobs(config.template_bucket(), prefix=f"{prefix}/" if prefix else None):
            name = blob.name.rsplit("/", 1)[-1]
            if name.lower().endswith(TEMPLATE_EXTS):
                out.append(Template(id=f"gcs:{blob.name}", name=Path(name).stem, source="gcs",
                                    is_default=Path(name).stem.lower() == "default", meta={"object": blob.name}))
        return out

    def _local_templates(self) -> list[Template]:
        if not self.local_dir or not self.local_dir.is_dir():
            return []
        return [Template(id=f"local:{p.name}", name=p.stem, source="local", is_default=p.stem.lower() == "default",
                         meta={"path": str(p)})
                for p in sorted(self.local_dir.iterdir())
                if p.suffix.lower() in TEMPLATE_EXTS and p.stem.lower() != "readme"]

    def list(self) -> list[Template]:
        user = self._firestore_templates() + self._gcs_templates() + self._local_templates()
        return user + [bundled_default()]

    # -- loading -------------------------------------------------------------------
    def _load(self, t: Template) -> Template:
        if t.content:
            return t
        if t.source == "firestore":
            if t.meta.get("content"):
                t.content = str(t.meta["content"]).strip()
            elif t.meta.get("gcs_uri") and self.gcs is not None:
                bucket, _, name = t.meta["gcs_uri"][5:].partition("/")
                t.content = _decode(name, self.gcs.bucket(bucket).blob(name).download_as_bytes())
        elif t.source == "gcs":
            name = t.meta["object"]
            t.content = _decode(name, self.gcs.bucket(config.template_bucket()).blob(name).download_as_bytes())
        elif t.source == "local":
            p = Path(t.meta["path"])
            t.content = _decode(p.name, p.read_bytes())
        if not t.content:
            raise TemplateNotFound(f"Template '{t.name}' has no readable content")
        return t

    def get(self, template_id: Optional[str] = None) -> Template:
        """By id or name; with no id: the user's default template, else the bundled default."""
        if not template_id:
            if self.source != "default":
                for t in self.list():
                    if t.is_default and t.source != "bundled":
                        return self._load(t)
            return bundled_default()
        if template_id == "default" or self.source == "default":
            return bundled_default()
        wanted = template_id.strip().lower()
        for t in self.list():
            if t.id.lower() == wanted or t.name.lower() == wanted:
                return self._load(t)
        raise TemplateNotFound(f"No template '{template_id}'. Available: {[t.name for t in self.list()]}")
